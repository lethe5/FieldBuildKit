"""D-95 canonical taxonomy workbook ingestion and lookup helpers.

This module is deliberately independent from the historical CSV/NIBR pipeline.  The workbook is
read with ``data_only=True`` so formulas and macros are never evaluated, and all validation is
completed before rows are returned to a caller.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

CANONICAL_FILENAME = "Rpt_2026-08-29_List.xlsx"
CANONICAL_URL_PREFIX = "https://species.nibr.go.kr/species-detail/"
CANONICAL_SHEET = "Data Sheet"
HEADER_ROWS = [1, 2]
DATA_START_ROW = 3
SCHEMA_REVISION = "D-95-1"
PIPELINE_REVISION = "D-95-canonical-workbook-1"
ALLOWED_SOURCE_KINDS = frozenset({"bundled_candidate", "user_upload"})

RANKS = ("Phylum", "Class", "Order", "Family", "Genus")
_HEADER_RANKS = RANKS + ("SubGenus", "Species")
SOURCE_COLUMNS = (
    "No", "관리분류군", "정이명여부", "학명", "대표국명", "명명자", "명명년도",
    "Phylum", "Phylum 국명", "Class", "Class 국명", "Order", "Order 국명",
    "Family", "Family 국명", "Genus", "Genus 국명", "SubGenus", "SubGenus 국명",
    "Species", "Species 국명", "콘텐츠 여부", "URL",
)
LOGICAL_COLUMNS = (
    "ktsn", "source_no", "taxon_status", "korean_name", "scientific_name",
    "scientific_name_normalized", "scientific_name_without_authority", "authority",
    "naming_year", "phylum_scientific_name", "phylum_korean_name",
    "class_scientific_name", "class_korean_name", "order_scientific_name", "order_korean_name",
    "family_scientific_name", "family_korean_name", "genus_scientific_name", "genus_korean_name",
    "accepted_ktsn", "source_url", "content_available", "source_row",
)

_AUTHORITY_MARKER_RE = re.compile(r"\s(?=[A-Z(])")
_SPACE_RE = re.compile(r"\s+")


class CanonicalReferenceError(ValueError):
    """A user-correctable canonical-source validation error."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _raw_text(value) -> str:
    """Return a source cell without changing its textual bytes/spacing.

    Canonical ``학명`` and ``URL`` are audit values.  Their comparison forms use
    :func:`_text`/the normalizer, but the values written to the reference tables
    must remain the original cell contents.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_canonical_scientific_name(value: str) -> dict:
    """Translate the normative R cleaner without replacing the raw source value."""
    x = _SPACE_RE.sub(" ", _text(value))
    x = re.sub(r"\s*\(\s*", " (", x)
    x = re.sub(r"\s*\)\s*", ") ", x)
    x = _SPACE_RE.sub(" ", x).strip()
    first_genus = re.sub(r"\s.*$", "", x)
    after_genus = re.sub(r"^[^ ]+\s+", "", x)
    marker = _AUTHORITY_MARKER_RE.search(after_genus)
    if marker is None:
        authority = ""
        scientific = x
    else:
        authority = after_genus[marker.end():].strip()
        scientific = (first_genus + " " + after_genus[:marker.start()]).strip()
    return {
        "scientific_name_normalized": x,
        "scientific_name_without_authority": scientific,
        "authority": authority,
    }


def extract_canonical_ktsn(url: str) -> dict:
    source_url = _raw_text(url)
    if not source_url.startswith(CANONICAL_URL_PREFIX):
        raise CanonicalReferenceError(
            "invalid_url_prefix",
            "URL이 허용된 국가생물종지식정보시스템 상세 URL 접두사와 일치하지 않습니다.",
        )
    ktsn = source_url[len(CANONICAL_URL_PREFIX):]
    if not ktsn.strip():
        raise CanonicalReferenceError("empty_ktsn", "URL에서 KTSN을 추출할 수 없습니다.")
    return {"source_url": source_url, "ktsn": ktsn}


def _error(code: str, message: str, *, provenance: dict | None = None) -> dict:
    result = {
        "success": False, "error_code": code, "error_message": message,
        "partial_rows": [], "promoted_project": False,
    }
    if provenance is not None:
        result["provenance"] = provenance
    return result


def _load_sheet(path: Path):
    if path.suffix.lower() != ".xlsx":
        raise CanonicalReferenceError("invalid_extension", "정식 참조 자료는 .xlsx 파일만 사용할 수 있습니다.")
    if not path.is_file():
        raise CanonicalReferenceError("missing_workbook", f"참조 엑셀 파일을 찾을 수 없습니다: {path}")
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exposes several different corruption exceptions
        raise CanonicalReferenceError("unreadable_workbook", f"참조 엑셀 파일을 읽을 수 없습니다: {path} ({exc})") from exc
    if CANONICAL_SHEET not in workbook.sheetnames:
        workbook.close()
        raise CanonicalReferenceError("missing_sheet", f"참조 엑셀 파일에 '{CANONICAL_SHEET}' 시트가 없습니다.")
    return workbook, workbook[CANONICAL_SHEET]


def _column_layout(sheet) -> list[str]:
    rows = list(sheet.iter_rows(min_row=1, max_row=2, values_only=True))
    if len(rows) < 2:
        raise CanonicalReferenceError("missing_header", "참조 엑셀 파일의 1·2행 헤더를 찾을 수 없습니다.")
    first, second = rows
    width = max(len(first), len(second))
    names: list[str] = []
    current_rank = None
    for index in range(width):
        top = _text(first[index] if index < len(first) else None)
        sub = _text(second[index] if index < len(second) else None)
        if top in _HEADER_RANKS:
            current_rank = top
        if sub == "국명" and current_rank:
            name = f"{current_rank} 국명"
        elif sub == "Name" and current_rank:
            name = current_rank
        else:
            name = top
        names.append(name)
    if names != list(SOURCE_COLUMNS):
        missing = [name for name in SOURCE_COLUMNS if name not in names]
        raise CanonicalReferenceError(
            "schema_mismatch",
            "참조 엑셀 헤더가 canonical schema와 일치하지 않습니다. "
            f"필수 열 누락: {missing or '열 순서/이름 확인 필요'}",
        )
    return names


def _materialize(path: Path, source_kind: str, expected_sha256: str | None = None) -> dict:
    initial_digest = hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else ""
    if expected_sha256:
        if initial_digest != expected_sha256:
            raise CanonicalReferenceError(
                "source_changed",
                "확인한 뒤 참조 엑셀 파일이 변경되었습니다. 파일을 다시 선택하고 확인해 주세요.",
            )
    workbook, sheet = _load_sheet(path)
    try:
        source_columns = _column_layout(sheet)
        indexes = {name: source_columns.index(name) for name in SOURCE_COLUMNS}
        rows: list[dict] = []
        seen_ktsn: set[str] = set()
        seen_urls: set[str] = set()
        accepted_ktsn: str | None = None
        taxon_groups: set[str] = set()
        accepted_count = synonym_count = 0
        for source_row, values in enumerate(sheet.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW):
            if not any(value is not None and _text(value) for value in values):
                continue
            raw = {column: (values[indexes[column]] if indexes[column] < len(values) else None) for column in SOURCE_COLUMNS}
            group = _text(raw["관리분류군"])
            if group != "관속식물류":
                raise CanonicalReferenceError("nonvascular_taxon", f"{source_row}행에 관속식물류가 아닌 분류군이 있습니다: {group or '빈 값'}")
            status = _text(raw["정이명여부"])
            if status not in {"정명", "이명"}:
                raise CanonicalReferenceError("invalid_taxon_status", f"{source_row}행의 정이명여부가 정명/이명이 아닙니다.")
            no = _text(raw["No"])
            if status == "정명":
                if not no or no == "-":
                    raise CanonicalReferenceError("accepted_no_missing", f"{source_row}행 정명에는 No가 필요합니다.")
            elif no != "-" or accepted_ktsn is None:
                raise CanonicalReferenceError("synonym_transition", f"{source_row}행 이명은 앞선 정명 뒤에 No='-'로 이어져야 합니다.")
            url_info = extract_canonical_ktsn(raw["URL"])
            ktsn, source_url = url_info["ktsn"], url_info["source_url"]
            if ktsn in seen_ktsn:
                raise CanonicalReferenceError("duplicate_ktsn", f"{source_row}행에서 중복 KTSN이 발견되었습니다: {ktsn}")
            if source_url in seen_urls:
                raise CanonicalReferenceError("duplicate_url", f"{source_row}행에서 중복 URL이 발견되었습니다.")
            seen_ktsn.add(ktsn); seen_urls.add(source_url)
            scientific_name = _raw_text(raw["학명"])
            if not _text(scientific_name):
                raise CanonicalReferenceError("scientific_name_missing", f"{source_row}행의 학명이 비어 있습니다.")
            normalized = normalize_canonical_scientific_name(scientific_name)
            if status == "정명":
                accepted_ktsn = ktsn; accepted_count += 1
            else:
                synonym_count += 1
            taxon_groups.add(group)
            logical = {
                "ktsn": ktsn, "source_no": no, "taxon_status": status,
                "korean_name": _text(raw["대표국명"]) or None,
                "scientific_name": scientific_name,
                **normalized,
                "naming_year": _text(raw["명명년도"]) or None,
                "accepted_ktsn": accepted_ktsn,
                "source_url": source_url,
                "content_available": _text(raw["콘텐츠 여부"]) or None,
                "source_row": source_row,
            }
            for rank in RANKS:
                logical[f"{rank.lower()}_scientific_name"] = _text(raw[rank]) or None
                logical[f"{rank.lower()}_korean_name"] = _text(raw[f"{rank} 국명"]) or None
            rows.append(logical)
        if not rows:
            raise CanonicalReferenceError("empty_workbook", "참조 엑셀 파일에 데이터 행이 없습니다.")
        final_digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if final_digest != initial_digest:
            raise CanonicalReferenceError(
                "source_changed",
                "참조 엑셀 파일을 읽는 동안 파일이 변경되었습니다. 파일을 다시 선택하고 확인해 주세요.",
            )
        digest = final_digest
        provenance = {
            "source_filename": path.name, "source_kind": source_kind,
            "sha256": digest, "byte_size": path.stat().st_size,
            "selection_timestamp": datetime.now(timezone.utc).isoformat(),
            "sheet_name": CANONICAL_SHEET, "header_rows": HEADER_ROWS,
            "data_start_row": DATA_START_ROW, "schema_revision": SCHEMA_REVISION,
            "pipeline_revision": PIPELINE_REVISION, "source_row_count": len(rows),
            "accepted_count": accepted_count, "synonym_count": synonym_count,
            "validation_result": "valid",
        }
        return {
            "success": True, "sheet_name": CANONICAL_SHEET, "header_rows": HEADER_ROWS,
            "data_start_row": DATA_START_ROW, "source_columns": source_columns,
            "logical_columns": list(LOGICAL_COLUMNS), "taxon_groups": sorted(taxon_groups),
            "row_count": len(rows), "accepted_count": accepted_count,
            "synonym_count": synonym_count, "rows": rows, "provenance": provenance,
            "source_kind": source_kind,
        }
    finally:
        workbook.close()


def ingest_canonical_workbook(workbook_path: str, source_kind: str = "bundled_candidate", expected_sha256: str | None = None) -> dict:
    path = Path(workbook_path)
    try:
        if source_kind not in ALLOWED_SOURCE_KINDS:
            allowed = ", ".join(sorted(ALLOWED_SOURCE_KINDS))
            raise CanonicalReferenceError(
                "invalid_source_kind",
                f"참조 자료 source_kind가 허용되지 않습니다: {source_kind!r}. "
                f"허용값은 {allowed}뿐입니다.",
            )
        result = _materialize(path, source_kind, expected_sha256)
        return result
    except CanonicalReferenceError as exc:
        return _error(exc.code, exc.message)
    except (OSError, PermissionError) as exc:
        return _error("source_io_error", f"참조 엑셀 파일에 접근할 수 없습니다: {path} ({exc})")
    except Exception as exc:  # keep the headless seam actionable and fail-closed
        return _error("source_validation_error", f"참조 엑셀 검증에 실패했습니다: {path} ({exc})")


def preview_canonical_workbook(workbook_path: str | Path, source_kind: str = "user_upload", sample_limit: int = 3) -> dict:
    """Validate one selected workbook, without discovering neighbouring files."""
    path = Path(workbook_path)
    result = ingest_canonical_workbook(str(path), source_kind=source_kind)
    entry = {
        "filename": path.name, "extension": path.suffix.lower(),
        # This is an internal candidate-selection handle, not provenance payload. The wizard uses
        # it to retain the discovered filesystem path without reconstructing it from a displayed
        # (possibly Unicode-normalized) filename.
        "path": str(path),
        "sheet_name": result.get("sheet_name"), "header_rows": result.get("header_rows", []),
        "sample_rows": [], "validation_status": "valid" if result.get("success") else "invalid",
        "source_kind": source_kind,
    }
    if result.get("success"):
        entry["provenance"] = dict(result["provenance"])
        workbook, sheet = _load_sheet(path)
        try:
            source_columns = list(result["source_columns"])
            entry["source_columns"] = source_columns
            entry["source_rows"] = [
                {
                    column: _raw_text(values[index]) if index < len(values) else ""
                    for index, column in enumerate(source_columns)
                }
                for values in sheet.iter_rows(
                    min_row=DATA_START_ROW,
                    max_row=DATA_START_ROW + max(0, int(sample_limit)) - 1,
                    values_only=True,
                )
                if any(value is not None and _text(value) for value in values)
            ]
        finally:
            workbook.close()
        if hashlib.sha256(path.read_bytes()).hexdigest() != result["provenance"]["sha256"]:
            entry.update(validation_status="invalid", error_code="source_changed",
                         error_message="미리보기를 읽는 동안 참조 파일이 변경되었습니다. 다시 선택해 주세요.")
            entry.pop("provenance", None)
            return entry
    for row in (result.get("rows") or [])[:max(0, int(sample_limit))]:
        entry["sample_rows"].append({
            "status": row["taxon_status"], "scientific_name": row["scientific_name"],
            "korean_name": row.get("korean_name"), "ktsn": str(row["ktsn"]),
            "phylum": row.get("phylum_scientific_name"), "class": row.get("class_scientific_name"),
            "order": row.get("order_scientific_name"), "family": row.get("family_scientific_name"),
            "genus": row.get("genus_scientific_name"),
        })
    if not result.get("success"):
        entry["error_code"] = result.get("error_code")
        entry["error_message"] = result.get("error_message")
    return entry


def inspect_ktsn_source_candidates(reference_tables_dir: str, upload_path: str | None = None, confirmed_path: str | None = None, sample_limit: int = 3) -> dict:
    directory = Path(reference_tables_dir)
    upload_resolved = Path(upload_path).resolve() if upload_path else None
    candidate_paths = [
        path for path in sorted(directory.glob("*.xlsx"))
        if upload_resolved is None or path.resolve() != upload_resolved
    ] if directory.is_dir() else []
    candidates = [preview_canonical_workbook(path, "bundled_candidate", sample_limit) for path in candidate_paths]
    # macOS may expose the same filename in decomposed Unicode while the API contract and UI use
    # NFC. Keep the actual paths separately so confirmation never reconstructs a path from a
    # display-normalized filename.
    for entry in candidates:
        entry["filename"] = unicodedata.normalize("NFC", entry["filename"])
    recommended = CANONICAL_FILENAME if any(item["filename"] == CANONICAL_FILENAME for item in candidates) else None
    upload = preview_canonical_workbook(upload_path, "user_upload", sample_limit) if upload_path else None
    selected = None
    if confirmed_path:
        confirmed = Path(confirmed_path)
        kind = "user_upload" if upload_path and confirmed.resolve() == Path(upload_path).resolve() else "bundled_candidate"
        chosen = upload if kind == "user_upload" else next(
            (item for item, source in zip(candidates, candidate_paths)
             if source.resolve() == confirmed.resolve()), None
        )
        if chosen and chosen.get("validation_status") == "valid":
            selected = {**chosen, "path": str(confirmed)}
    result = {
        "recommended_filename": recommended, "selected": selected, "candidates": candidates,
        "can_continue": selected is not None,
    }
    if upload is not None:
        result["upload"] = upload
    if not candidates and upload is None:
        result.update(error_code="no_candidate", error_message="사용할 수 있는 .xlsx 참조 자료가 없습니다. 올바른 엑셀 파일을 업로드해 주세요.")
    elif selected is None and confirmed_path:
        result.update(error_code="source_not_confirmed", error_message="미리보기를 확인한 유효한 .xlsx 파일을 선택해 주세요.")
    return result


def match_canonical_ktsn(scientific_name_without_authority: str, rows: list[dict]) -> dict:
    result = {
        "matched": False, "matched_ktsn": None, "taxon_status": None,
        "accepted_ktsn": None, "selected_korean_name": None,
        "selected_scientific_name": None, "selected_scientific_name_without_authority": None,
        "ambiguous": False,
    }
    target = _text(scientific_name_without_authority)
    matches = [row for row in rows if _text(row.get("scientific_name_without_authority")) == target]
    if not matches:
        return result
    accepted_groups = {_text(row.get("accepted_ktsn")) for row in matches}
    if len(accepted_groups) > 1:
        result.update(matched=True, ambiguous=True)
        return result
    direct = matches[0]
    accepted_ktsn = _text(direct.get("accepted_ktsn"))
    accepted = next((row for row in rows if _text(row.get("ktsn")) == accepted_ktsn), None)
    result.update(matched=True, matched_ktsn=direct.get("ktsn"), taxon_status=direct.get("taxon_status"), accepted_ktsn=accepted_ktsn)
    if accepted:
        result.update(selected_korean_name=accepted.get("korean_name"), selected_scientific_name=accepted.get("scientific_name"), selected_scientific_name_without_authority=accepted.get("scientific_name_without_authority"))
    return result


def aggregate_taxonomy_report(observations: list[dict], taxonomy_rows: list[dict] | None, survey_type: str) -> dict:
    ordinary = list(observations or [])
    if survey_type == "vegetation_mapping":
        return {"ordinary_observations": ordinary, "aggregations": {}, "taxonomy_applicability": "not_applicable", "limitations": []}
    if taxonomy_rows is None:
        return {"ordinary_observations": ordinary, "aggregations": {}, "limitations": ["taxonomy_reference_unavailable"]}
    by_ktsn = {_text(row.get("ktsn")): row for row in taxonomy_rows if _text(row.get("ktsn"))}
    aggregations: dict[str, dict[tuple[str, str], dict]] = {rank: {} for rank in RANKS}
    limitations: list[str] = []

    def add_limitation(message: str) -> None:
        if message not in limitations:
            limitations.append(message)

    for observation in ordinary:
        ktsn = _text(observation.get("selected_ktsn"))
        if not ktsn:
            add_limitation("blank_ktsn")
            continue
        row = by_ktsn.get(ktsn)
        if row is None:
            # Keep the historical raw-KTSN limitation shape for callers that use it as a
            # machine-readable missing-key list; the report renderer supplies the Korean context.
            add_limitation(ktsn)
            continue
        accepted = by_ktsn.get(_text(row.get("accepted_ktsn"))) if row.get("accepted_ktsn") else row
        if accepted is None:
            add_limitation(ktsn)
            continue
        for rank in RANKS:
            scientific = _text(accepted.get(f"{rank.lower()}_scientific_name"))
            korean = _text(accepted.get(f"{rank.lower()}_korean_name"))
            if not scientific or not korean:
                add_limitation(
                    f"KTSN {ktsn}의 {rank} 계층은 학명/국명 누락으로 분류 집계에서 제외되었습니다"
                )
                continue
            key = (scientific, korean)
            bucket = aggregations[rank].setdefault(key, {"scientific_name": scientific, "korean_name": korean, "observation_ids": set(), "ktsns": set()})
            bucket["observation_ids"].add(_text(observation.get("observation_id")))
            bucket["ktsns"].add(ktsn)
    serial = {}
    for rank, buckets in aggregations.items():
        serial[rank] = [{"scientific_name": b["scientific_name"], "korean_name": b["korean_name"], "observation_count": len(b["observation_ids"]), "distinct_ktsn_count": len(b["ktsns"])} for b in buckets.values()]
    return {"ordinary_observations": ordinary, "aggregations": serial, "limitations": limitations}
