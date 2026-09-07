"""Section 13.3 reference-asset validation and bundling (post-MVP guaranteed-manual baseline).

The D-95 canonical workbook/release path is the normal build contract.  The CSV/older-workbook
helpers later in this module are retained only for the explicitly selected legacy compatibility
mode and are not consulted by canonical project creation or release staging.

- The normal D-95 build/release path validates and materializes the canonical workbook through
  :func:`prepare_filtered_reference_bundle`; it does not require the older source tables.
- FR-QPB-105 (further revised; Decision Log D-70) and the related CSV pipeline remain below as
  explicit legacy-compatibility helpers for pre-D-95 projects and fixture coverage only.
- FR-QPB-112 (revised; Decision Log D-47; original Decision Log D-35): bundle a build-time-
  extracted, 5-column KTSN lookup file (see FR-QPB-118) -- never the complete, unmodified raw
  KTSN CSV -- plus the complete, unmodified probability-raster set (this raster half of the rule
  is unchanged from Decision Log D-35), into the generated project's own ``reference/`` folder.
- FR-QPB-113 (Decision Log D-41): extract only the ``KTSN`` column from the NIBR accepted-taxon
  xlsx into a small, single-purpose lookup-list file, bundled the same way.
- FR-QPB-118 (further revised; Decision Log D-48, complete five-step form): before the existing
  5-column extraction (now Step 5) runs, four ordered row-filtering steps narrow the complete raw
  KTSN CSV down to only the rows eligible for on-device matching:

  1. Plantae-kingdom subtree filter -- keep only rows that are transitive descendants of
     ``ktsn=120000098391`` via the ``p_ktsn`` parent-pointer chain.
  2. Rank filter -- keep only rows whose ``rank_id`` (parsed as an integer) is ``>= 700``
     (species-level-and-below); a missing/non-numeric ``rank_id`` excludes the row.
  3. Vascular-plants-only phylum allowlist -- keep only rows whose ``r200_nm`` is exactly one of
     ``Magnoliophyta``, ``Pteridophyta``, ``Pinophyta``, ``Filicophyta``, ``Lycopodiophyta``,
     ``Psilophyta``, ``Sphenophyta``; a missing/blank ``r200_nm`` excludes the row.
  4. Duplicate-scientific-name canonical-KTSN resolution -- group the rows surviving Step 3 by
     normalized ``taxon_full_nm`` (the same ``<em>``/``</em>``-tag-stripping-plus-whitespace-
     collapsing normalization :mod:`qfield_builder.ktsn_match` already implements); for a group of
     more than one row, resolve via the NIBR xlsx's ``관속식물류`` sheet (``학명``/``KTSN``
     columns) -- keep exactly one row iff the sheet yields exactly one distinct ``KTSN`` value that
     also equals one of the group's own ``ktsn`` values; drop the entire group otherwise.

  All five steps run once, at build time, entirely in pure Python, never at runtime on-device.

- DR-QPB-072/FR-QPB-124-127 (Decision Log D-66/D-67/D-68): a real, indexed GeoPackage table (the
  "KTSN accepted-name lookup table"), bundled inside the generated project's own survey `.gpkg`
  (never a second `.gpkg` file), containing only the currently accepted (non-synonym) taxa --
  `derive_accepted_name_lookup_table()` below reuses this same module's shared Steps 1-4
  computation, then applies its own new Step 5 (`taxon_jm_nm == '정명'` only). Unlike every other
  bundling rule in this module, this one is bundled unconditionally for every Types 1-3 project
  (FR-QPB-127), independent of whether the identification subsystem is enabled. The actual
  GeoPackage table creation/indexing/insertion is performed by
  :func:`qfield_builder.gpkg.add_ktsn_lookup_table`, not by this module, which only derives the
  row data.

All paths this module writes into the generated project are relative to the project folder
(FR-QPB-091/DR-QPB-012's existing no-absolute-paths discipline, applied here per
FR-QPB-112/113/118).
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
from pathlib import Path

from . import canonical_reference
from .errors import ReferenceDataInvalidError, ReferenceDataMissingError
from .ktsn_match import _strip_em_tags as _normalize_taxon_full_nm

KTSN_CSV_RELATIVE_SUBPATH = Path("tables") / "tb_leco_nib_ktsn_dtl_gat.csv"
RASTER_DIR_RELATIVE_SUBPATH = Path("rasters") / "bce_inverse_corrected_probability_maps"
NATIONAL_LIST_XLSX_RELATIVE_SUBPATH = Path("tables") / "2025년 국가생물종목록_v1.0.xlsx"
NATIONAL_LIST_XLSX_SHEET_NAME = "62,604종"
NATIONAL_LIST_XLSX_KTSN_COLUMN_HEADER_ROW = 1
NATIONAL_LIST_XLSX_DATA_START_ROW = 3

REQUIRED_KTSN_CSV_COLUMNS = (
    "ktsn",
    "taxon_full_nm",
    "taxon_kor_nm",
    "taxon_jm_nm",
    "correct_list",
)

# --- FR-QPB-118 (further revised; Decision Log D-48) five-step pipeline constants ----------------

# Step 1: the Plantae kingdom's own `ktsn` (this row does not itself appear in the species-detail
# table -- only rows reachable from it via the `p_ktsn` chain are "surviving").
PLANTAE_ROOT_KTSN = "120000098391"

# Step 2: species-level-and-below (infraspecific) rank codes only.
STEP2_MIN_RANK_ID = 700

# Step 3: the seven vascular-plant phyla (angiosperms, ferns/fern-allies, gymnosperms) among the
# phyla the Plantae-kingdom subtree contains -- excludes the three bryophyte phyla and the three
# algae-related phyla this dataset classifies under Plantae (Decision Log D-48).
VASCULAR_PHYLUM_ALLOWLIST = (
    "Magnoliophyta",
    "Pteridophyta",
    "Pinophyta",
    "Filicophyta",
    "Lycopodiophyta",
    "Psilophyta",
    "Sphenophyta",
)

# Step 4: the NIBR xlsx workbook's `관속식물류` (vascular plants) sheet -- a different sheet of the
# same workbook already required by FR-QPB-113/Decision Log D-40/D-41 (which reads `62,604종`).
NIBR_VASCULAR_SHEET_NAME = "관속식물류"
NIBR_VASCULAR_SHEET_NAME_HEADER = "학명"
NIBR_VASCULAR_SHEET_KTSN_HEADER = "KTSN"
NIBR_VASCULAR_SHEET_DATA_START_ROW = 3

# Project-relative destination paths (FR-QPB-090's existing `reference/` folder anticipation).
PROJECT_REFERENCE_DIR = "reference"
# FR-QPB-112 (revised)/FR-QPB-118 (Decision Log D-47/D-48): the build-time-extracted, 5-column
# lookup file, produced from only the rows surviving all five ordered pipeline steps. This
# replaces the former `reference/tables/tb_leco_nib_ktsn_dtl_gat.csv` complete-raw-CSV bundling
# location, which must no longer be created at all.
PROJECT_KTSN_LOOKUP_RELPATH = f"{PROJECT_REFERENCE_DIR}/ktsn_lookup.csv"
PROJECT_RASTER_DIR_RELPATH = (
    f"{PROJECT_REFERENCE_DIR}/rasters/bce_inverse_corrected_probability_maps"
)
PROJECT_NATIONAL_LIST_RELPATH = f"{PROJECT_REFERENCE_DIR}/nibr_accepted_ktsn_list.txt"

# D-89: distributed applications receive this validated, generated subset rather than the raw
# tables directory.  The manifest deliberately lives one level above it so its own digest can be
# checked without an impossible self-referential hash.
FILTERED_REFERENCE_RELPATH = Path("filtered")
FILTERED_KTSN_LOOKUP_NAME = "ktsn_lookup.csv"
FILTERED_ACCEPTED_LOOKUP_NAME = "accepted_name_lookup.csv"
FILTERED_NIBR_LIST_NAME = "nibr_accepted_ktsn_list.txt"
FILTERED_MANIFEST_NAME = "bundle_manifest.json"
FILTERED_EXPECTED_KTSN_ROWS = 20_914
FILTERED_PIPELINE_REVISION = "FR-QPB-118-D48"
CANONICAL_FILTERED_PIPELINE_REVISION = "D-95-canonical-release-1"
CANONICAL_WORKBOOK_RELATIVE_SUBPATH = Path("tables") / canonical_reference.CANONICAL_FILENAME
RELEASE_SOURCE_MANIFEST_NAME = "canonical_source_manifest.json"

# --- DR-QPB-072/FR-QPB-124-127 (Decision Log D-66/D-67/D-68) -- the bundled KTSN accepted-name ---
# --- lookup table (a real, indexed GeoPackage table, not a project-relative file) -----------------

# Decision Log D-4/15b's existing accepted/canonical nomenclatural-status flag: a row is the
# accepted (non-synonym) record for its scientific name iff `taxon_jm_nm == ACCEPTED_TAXON_STATUS_
# VALUE` exactly (FR-QPB-107, further revised; restated by FR-QPB-126's own Step 5).
ACCEPTED_TAXON_STATUS_VALUE = "정명"

# The 3 columns FR-QPB-126/HARNESS_CONTRACT.md function 17 require, at minimum, for the bundled
# accepted-name lookup table.
ACCEPTED_NAME_LOOKUP_TABLE_COLUMNS = ("ktsn", "taxon_kor_nm", "taxon_full_nm")

# DR-QPB-072/HARNESS_CONTRACT.md ("build_project config/return additions"): the specification does
# not mandate a specific literal GeoPackage table name for the bundled accepted-name lookup table
# -- this is this implementation's own choice, reported back via `build_project()`'s
# `ktsn_lookup_table_name` return key.
KTSN_LOOKUP_TABLE_NAME = "ktsn_accepted_name_lookup"


def _validate_ktsn_csv(csv_path: Path) -> None:
    if not csv_path.is_file():
        raise ReferenceDataMissingError(
            f"식별 기능에 필요한 KTSN 참조 CSV 파일을 찾을 수 없습니다: {csv_path}. "
            "'storage/reference/tables/tb_leco_nib_ktsn_dtl_gat.csv' 파일이 올바른 위치에 "
            "있는지 확인한 뒤 다시 시도해 주세요."
        )
    try:
        with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            header = next(reader, None)
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ReferenceDataInvalidError(
            f"KTSN 참조 CSV 파일을 읽는 중 오류가 발생했습니다: {csv_path} ({exc})"
        ) from exc

    if header is None:
        raise ReferenceDataInvalidError(
            f"KTSN 참조 CSV 파일이 비어 있습니다 (헤더 행이 없습니다): {csv_path}"
        )
    missing = [col for col in REQUIRED_KTSN_CSV_COLUMNS if col not in header]
    if missing:
        raise ReferenceDataInvalidError(
            f"KTSN 참조 CSV 파일에 필수 열이 없습니다: {missing} (파일: {csv_path}). "
            f"필수 열: {list(REQUIRED_KTSN_CSV_COLUMNS)}"
        )


def _validate_raster_dir(raster_dir: Path) -> None:
    if not raster_dir.is_dir():
        raise ReferenceDataMissingError(
            f"식별 기능에 필요한 확률 래스터 폴더를 찾을 수 없습니다: {raster_dir}. "
            "'storage/reference/rasters/bce_inverse_corrected_probability_maps/' 폴더가 올바른 "
            "위치에 있는지 확인한 뒤 다시 시도해 주세요."
        )
    entries = list(raster_dir.iterdir())
    unexpected = sorted(
        path.name for path in entries if path.is_file() and path.suffix.lower() != ".tif"
    )
    if unexpected:
        raise ReferenceDataInvalidError(
            f"확률 래스터 폴더에 TIFF가 아닌 파일이 있습니다: {unexpected}"
        )
    nested = sorted(path.name for path in entries if path.is_dir())
    if nested:
        raise ReferenceDataInvalidError(
            f"확률 래스터 폴더에 허용되지 않는 하위 폴더가 있습니다: {nested}"
        )
    tif_files = sorted(path for path in entries if path.is_file() and path.suffix.lower() == ".tif")
    if not tif_files:
        raise ReferenceDataMissingError(
            f"확률 래스터 폴더에 TIFF 파일이 없습니다: {raster_dir}. "
            "필요한 관속식물 확률 래스터 세트를 확인한 뒤 다시 시도해 주세요."
        )
    empty = [path.name for path in tif_files if not path.is_file() or path.stat().st_size == 0]
    if empty:
        raise ReferenceDataInvalidError(
            f"확률 래스터에 비어 있거나 읽을 수 없는 TIFF 파일이 있습니다: {empty}"
        )


def validate_reference_data(reference_data_dir: str) -> tuple[Path, Path]:
    """Validate the pre-D-95 source contract for explicit legacy compatibility callers only."""
    base = Path(reference_data_dir)
    # Prefer raw inputs whenever they exist.  A stale `filtered/` directory beside changed raw
    # tables must never shadow those inputs and be reused as if it were current.
    raw_csv = base / KTSN_CSV_RELATIVE_SUBPATH
    if not raw_csv.is_file() and (base / FILTERED_REFERENCE_RELPATH).is_dir():
        return validate_filtered_reference_data(base)
    csv_path = base / KTSN_CSV_RELATIVE_SUBPATH
    raster_dir = base / RASTER_DIR_RELATIVE_SUBPATH
    _validate_ktsn_csv(csv_path)
    _validate_raster_dir(raster_dir)
    return csv_path, raster_dir


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_filtered_csv(path: Path, expected_header: tuple[str, ...], label: str) -> list[list[str]]:
    """Read a filtered CSV with an exact header and rectangular rows."""
    try:
        with path.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            header = next(reader, None)
            rows = list(reader)
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ReferenceDataInvalidError(
            f"필터링된 {label}을 읽을 수 없습니다: {path} ({exc})"
        ) from exc
    if header != list(expected_header):
        raise ReferenceDataInvalidError(f"필터링된 {label}의 열이 올바르지 않습니다: {path}")
    width = len(expected_header)
    if any(len(row) != width for row in rows):
        raise ReferenceDataInvalidError(f"필터링된 {label}에 열 수가 다른 행이 있습니다: {path}")
    return rows


def _validate_artifact_hash(entry: dict, path: Path, name: str) -> None:
    digest = entry.get("sha256")
    if not isinstance(digest, str) or len(digest) != 64:
        raise ReferenceDataInvalidError(f"manifest의 자료 해시가 올바르지 않습니다: {name}")
    if digest != _sha256_file(path):
        raise ReferenceDataInvalidError(f"필터링된 참조 자료의 해시가 manifest와 다릅니다: {name}")


def _filtered_manifest_path(base: Path) -> Path:
    for candidate in (
        base / FILTERED_MANIFEST_NAME,
        base / FILTERED_REFERENCE_RELPATH / FILTERED_MANIFEST_NAME,
    ):
        if candidate.is_file():
            return candidate
    raise ReferenceDataMissingError(
        f"필터링된 참조 번들의 manifest 파일을 찾을 수 없습니다: {base / FILTERED_MANIFEST_NAME}"
    )


def _validate_filtered_manifest(
    base: Path, *, require_release_contract: bool = False
) -> tuple[Path, Path, Path, Path]:
    """Validate a D-89 bundle and return lookup, accepted, NIBR and raster paths.

    The ordinary project test seam retains compatibility with the original compact fixture
    manifest.  Release staging passes ``require_release_contract=True`` and therefore requires
    every D-89 provenance/inventory field with no compatibility defaults.
    """
    filtered = base / FILTERED_REFERENCE_RELPATH
    if not filtered.is_dir():
        raise ReferenceDataMissingError(f"필터링된 참조 자료 폴더를 찾을 수 없습니다: {filtered}")
    manifest_path = _filtered_manifest_path(base)
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ReferenceDataInvalidError(
            f"필터링된 참조 번들의 manifest를 읽을 수 없습니다: {manifest_path} ({exc})"
        ) from exc
    artifacts = manifest.get("artifacts") if isinstance(manifest, dict) else None
    if not isinstance(artifacts, dict):
        raise ReferenceDataInvalidError(
            f"필터링된 참조 번들의 manifest 형식이 잘못되었습니다: {manifest_path}"
        )

    canonical_release = manifest.get("pipeline_revision") == CANONICAL_FILTERED_PIPELINE_REVISION
    if require_release_contract:
        expected_pipeline = CANONICAL_FILTERED_PIPELINE_REVISION
        if manifest.get("pipeline_revision") != expected_pipeline:
            raise ReferenceDataInvalidError(
                "release filtered bundle의 pipeline_revision이 D-95 canonical release와 "
                "일치하지 않습니다"
            )
        provenance = manifest.get("source_provenance")
        if not isinstance(provenance, dict) or provenance.get("input_kind") != (
            "canonical_workbook"
        ):
            raise ReferenceDataInvalidError(
                "release filtered bundle에 canonical workbook provenance가 없습니다"
            )
        expected_source_paths = {
            "canonical_workbook": CANONICAL_WORKBOOK_RELATIVE_SUBPATH,
        }
        for source_name, expected_source_path in expected_source_paths.items():
            source_entry = provenance.get(source_name)
            if not isinstance(source_entry, dict):
                raise ReferenceDataInvalidError(
                    f"release manifest에 source provenance가 없습니다: {source_name}"
                )
            source_path = Path(str(source_entry.get("path") or ""))
            if source_path.is_absolute() or ".." in source_path.parts or not source_path.parts:
                raise ReferenceDataInvalidError(
                    f"release manifest의 source provenance 경로가 안전하지 않습니다: {source_name}"
                )
            if source_path != expected_source_path:
                raise ReferenceDataInvalidError(
                    f"release manifest의 source provenance 경로가 올바르지 않습니다: {source_name}"
                )
            digest = source_entry.get("sha256")
            if not isinstance(digest, str) or len(digest) != 64:
                raise ReferenceDataInvalidError(
                    f"release manifest의 source provenance 해시가 올바르지 않습니다: {source_name}"
                )
        if set(artifacts) != {
            FILTERED_KTSN_LOOKUP_NAME,
            FILTERED_ACCEPTED_LOOKUP_NAME,
            FILTERED_NIBR_LIST_NAME,
        }:
            raise ReferenceDataInvalidError(
                "release manifest의 filtered artifact 목록이 정확히 일치하지 않습니다"
            )

    expected = {
        FILTERED_KTSN_LOOKUP_NAME: ("csv",),
        FILTERED_ACCEPTED_LOOKUP_NAME: ("csv",),
        FILTERED_NIBR_LIST_NAME: ("txt",),
    }
    resolved: dict[str, Path] = {}
    for name, suffixes in expected.items():
        entry = artifacts.get(name)
        if not isinstance(entry, dict):
            raise ReferenceDataInvalidError(f"manifest에 필수 필터 자료가 없습니다: {name}")
        default_path = "" if require_release_contract else f"filtered/{name}"
        relative = Path(str(entry.get("path") or default_path))
        if relative.is_absolute() or ".." in relative.parts:
            raise ReferenceDataInvalidError(f"manifest의 자료 경로가 안전하지 않습니다: {name}")
        path = base / relative
        if (
            path.parent != filtered
            or path.suffix.lower().lstrip(".") not in suffixes
            or not path.is_file()
        ):
            raise ReferenceDataInvalidError(
                f"필터링된 참조 자료가 없거나 경로가 잘못되었습니다: {name}"
            )
        if require_release_contract and relative != FILTERED_REFERENCE_RELPATH / name:
            raise ReferenceDataInvalidError(
                f"release manifest의 artifact 경로가 올바르지 않습니다: {name}"
            )
        if require_release_contract:
            row_count = entry.get("row_count")
            if isinstance(row_count, bool) or not isinstance(row_count, int) or row_count < 0:
                raise ReferenceDataInvalidError(
                    f"release manifest의 artifact 행 수가 올바르지 않습니다: {name}"
                )
        _validate_artifact_hash(entry, path, name)
        resolved[name] = path

    if require_release_contract:
        expected_paths = set(resolved.values())
        extra_files = sorted(
            path.relative_to(filtered).as_posix()
            for path in filtered.rglob("*")
            if path.is_file() and path not in expected_paths
        )
        if extra_files:
            raise ReferenceDataInvalidError(
                f"release filtered bundle에 manifest에 없는 파일이 있습니다: {extra_files}"
            )

    lookup = resolved[FILTERED_KTSN_LOOKUP_NAME]
    rows = _read_filtered_csv(lookup, REQUIRED_KTSN_CSV_COLUMNS, "KTSN lookup")
    row_count = artifacts[FILTERED_KTSN_LOOKUP_NAME].get("row_count")
    expected_row_count = (
        row_count
        if require_release_contract or canonical_release
        else FILTERED_EXPECTED_KTSN_ROWS
    )
    if row_count != len(rows) or (
        not require_release_contract
        and not canonical_release
        and len(rows) != FILTERED_EXPECTED_KTSN_ROWS
    ):
        raise ReferenceDataInvalidError(
            f"필터링된 KTSN lookup 행 수가 올바르지 않습니다: {len(rows)} "
            f"(expected {expected_row_count})"
        )

    lookup_by_ktsn: dict[str, list[str]] = {}
    for row_number, row in enumerate(rows, start=2):
        ktsn, scientific_name, _korean_name, status, correct_list = row
        if not ktsn.strip() or not scientific_name.strip():
            raise ReferenceDataInvalidError(
                f"필터링된 KTSN lookup에 필수 값이 없는 행이 있습니다: {lookup}:{row_number}"
            )
        if ktsn.strip() in lookup_by_ktsn:
            raise ReferenceDataInvalidError(
                f"필터링된 KTSN lookup에 중복 KTSN이 있습니다: {lookup}:{row_number}"
            )
        # D-89 staging must never ship a partially parseable transition column.  Empty source
        # values are canonicalized to [] when the raw source is staged; an empty cell in an
        # already-filtered input therefore remains invalid rather than being silently repaired.
        try:
            parsed_correct_list = json.loads(correct_list)
        except (TypeError, json.JSONDecodeError) as exc:
            raise ReferenceDataInvalidError(
                f"필터링된 KTSN lookup의 correct_list가 JSON 배열이 아닙니다: "
                f"{lookup}:{row_number} ({exc})"
            ) from exc
        if not isinstance(parsed_correct_list, list):
            raise ReferenceDataInvalidError(
                f"필터링된 KTSN lookup의 correct_list가 JSON 배열이 아닙니다: "
                f"{lookup}:{row_number}"
            )
        lookup_by_ktsn[ktsn.strip()] = [status]

    accepted = resolved[FILTERED_ACCEPTED_LOOKUP_NAME]
    accepted_rows = _read_filtered_csv(
        accepted, ACCEPTED_NAME_LOOKUP_TABLE_COLUMNS, "accepted-name lookup"
    )
    accepted_count = artifacts[FILTERED_ACCEPTED_LOOKUP_NAME].get("row_count")
    if accepted_count != len(accepted_rows) or not accepted_rows:
        raise ReferenceDataInvalidError(
            f"필터링된 accepted-name lookup 행 수가 올바르지 않습니다: {len(accepted_rows)}"
        )
    accepted_ktsns: set[str] = set()
    for row_number, row in enumerate(accepted_rows, start=2):
        ktsn, _korean_name, scientific_name = row
        normalized_ktsn = ktsn.strip()
        if not normalized_ktsn or not scientific_name.strip():
            raise ReferenceDataInvalidError(
                f"필터링된 accepted-name lookup에 필수 값이 없는 행이 있습니다: "
                f"{accepted}:{row_number}"
            )
        if normalized_ktsn in accepted_ktsns:
            raise ReferenceDataInvalidError(
                f"필터링된 accepted-name lookup에 중복 KTSN이 있습니다: {accepted}:{row_number}"
            )
        accepted_ktsns.add(normalized_ktsn)
        lookup_status = lookup_by_ktsn.get(normalized_ktsn)
        if lookup_status is None or lookup_status[0] != ACCEPTED_TAXON_STATUS_VALUE:
            raise ReferenceDataInvalidError(
                f"accepted-name lookup의 KTSN이 정명 KTSN lookup과 일치하지 않습니다: "
                f"{accepted}:{row_number}"
            )

    nibr = resolved[FILTERED_NIBR_LIST_NAME]
    try:
        nibr_lines = nibr.read_text(encoding="utf-8-sig").splitlines()
    except (OSError, UnicodeDecodeError) as exc:
        raise ReferenceDataInvalidError(
            f"필터링된 NIBR 목록을 읽을 수 없습니다: {nibr} ({exc})"
        ) from exc
    if not nibr_lines or any(
        not line
        or line != line.strip()
        or any(ch.isspace() for ch in line)
        or "," in line
        or "\t" in line
        for line in nibr_lines
    ):
        raise ReferenceDataInvalidError(
            f"필터링된 NIBR 목록은 빈 줄이나 공백을 포함하지 않는 한 열 형식이어야 합니다: {nibr}"
        )
    if artifacts[FILTERED_NIBR_LIST_NAME].get("row_count") != len(nibr_lines):
        raise ReferenceDataInvalidError(f"필터링된 NIBR 목록 행 수가 manifest와 다릅니다: {nibr}")

    for path in filtered.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}:
            raise ReferenceDataInvalidError(
                f"필터링된 번들에 원본 스프레드시트가 포함되어 있습니다: {path}"
            )
    raster_dir = base / RASTER_DIR_RELATIVE_SUBPATH
    _validate_raster_dir(raster_dir)
    raster_entries = manifest.get("raster_files")
    if require_release_contract and not isinstance(raster_entries, dict):
        raise ReferenceDataInvalidError(
            "release manifest에 필수 확률 래스터 inventory가 없습니다"
        )
    if raster_entries is not None:
        actual_raster_paths = [
            path
            for path in raster_dir.iterdir()
            if path.is_file() and path.suffix.lower() == ".tif"
        ]
        unexpected_raster_files = [
            path.name
            for path in raster_dir.iterdir()
            if not path.is_file() or path.suffix.lower() != ".tif"
        ]
        if unexpected_raster_files:
            raise ReferenceDataInvalidError(
                "확률 래스터 폴더에 허용되지 않는 파일이 있습니다: "
                f"{sorted(unexpected_raster_files)}"
            )
        if not isinstance(raster_entries, dict) or set(raster_entries) != {
            path.name for path in actual_raster_paths
        }:
            raise ReferenceDataInvalidError(
                f"manifest의 확률 래스터 목록이 실제 TIFF 세트와 다릅니다: {raster_dir}"
            )
        for name, entry in raster_entries.items():
            if not isinstance(entry, dict):
                raise ReferenceDataInvalidError(
                    f"manifest의 래스터 항목이 올바르지 않습니다: {name}"
                )
            expected_path = raster_dir / name
            if Path(name).name != name or Path(name).suffix.lower() != ".tif":
                raise ReferenceDataInvalidError(
                    f"manifest의 래스터 파일명이 올바르지 않습니다: {name}"
                )
            if not expected_path.is_file() or expected_path.stat().st_size == 0:
                raise ReferenceDataInvalidError(f"manifest의 래스터 파일이 비어 있습니다: {name}")
            relative = Path(str(entry.get("path") or ""))
            if (
                relative.is_absolute()
                or ".." in relative.parts
                or relative != RASTER_DIR_RELATIVE_SUBPATH / name
                or relative.name != name
            ):
                raise ReferenceDataInvalidError(
                    f"manifest의 래스터 경로가 안전하지 않습니다: {name}"
                )
            if entry.get("size") != expected_path.stat().st_size:
                raise ReferenceDataInvalidError(f"manifest의 래스터 크기가 다릅니다: {name}")
            _validate_artifact_hash(entry, expected_path, name)
    return (
        lookup,
        resolved[FILTERED_ACCEPTED_LOOKUP_NAME],
        resolved[FILTERED_NIBR_LIST_NAME],
        raster_dir,
    )


def is_filtered_reference_data(reference_data_dir: str | Path) -> bool:
    base = Path(reference_data_dir)
    # A filtered directory is authoritative only when this is genuinely a filtered-only source.
    # If raw inputs are present, any neighboring filtered directory may be stale and must not be
    # selected by project generation or release staging.
    return not (base / KTSN_CSV_RELATIVE_SUBPATH).is_file() and (
        base / FILTERED_REFERENCE_RELPATH
    ).is_dir()


def validate_filtered_reference_data(reference_data_dir: str | Path) -> tuple[Path, Path]:
    lookup, _accepted, _nibr, raster_dir = _validate_filtered_manifest(Path(reference_data_dir))
    return lookup, raster_dir


# ================================================================================================
# FR-QPB-118 (further revised; Decision Log D-48) -- the five-step row-filtering/extraction pipeline
# ================================================================================================


def _load_full_csv_rows(csv_path: Path) -> list[dict]:
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _build_parent_map(rows: list[dict]) -> dict[str, str | None]:
    parent_of: dict[str, str | None] = {}
    for row in rows:
        ktsn = row.get("ktsn")
        if ktsn:
            parent_of[ktsn] = row.get("p_ktsn") or None
    return parent_of


def _resolves_to_plantae_root(
    start: str | None, parent_of: dict[str, str | None], status: dict[str, bool]
) -> bool:
    """Walks the `p_ktsn` parent-pointer chain starting at `start`, memoizing each visited node's
    final outcome in `status` so the overall Step 1 computation across all rows is amortized O(n).
    A chain that terminates (empty/blank/absent `p_ktsn`), that is broken (a `p_ktsn` value with no
    corresponding row `ktsn` in the same file), or that cycles back on itself without ever reaching
    `PLANTAE_ROOT_KTSN`, resolves to False."""
    path: list[str] = []
    current = start
    outcome = False
    while True:
        if not current:
            outcome = False
            break
        if current == PLANTAE_ROOT_KTSN:
            outcome = True
            break
        if current in status:
            outcome = status[current]
            break
        if current in path:
            outcome = False  # cycle: never reaches the root
            break
        path.append(current)
        current = parent_of.get(current)  # None if the chain is broken (parent row absent)
    for node in path:
        status[node] = outcome
    return outcome


def _step1_plantae_subtree_filter(rows: list[dict]) -> list[dict]:
    """FR-QPB-118 Step 1 (Decision Log D-48): keep only rows that are transitive descendants of
    `PLANTAE_ROOT_KTSN` via the `p_ktsn` parent-pointer chain."""
    parent_of = _build_parent_map(rows)
    status: dict[str, bool] = {}
    return [
        row
        for row in rows
        if _resolves_to_plantae_root(row.get("p_ktsn") or None, parent_of, status)
    ]


def _parse_rank_id(value) -> int | None:
    if value is None:
        return None
    text = value.strip() if isinstance(value, str) else str(value).strip()
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _step2_rank_filter(rows: list[dict]) -> list[dict]:
    """FR-QPB-118 Step 2 (Decision Log D-48): keep only rows whose `rank_id`, parsed as a base-10
    integer, is `>= STEP2_MIN_RANK_ID`. A missing/blank/non-numeric `rank_id` excludes the row."""
    survivors = []
    for row in rows:
        rank = _parse_rank_id(row.get("rank_id"))
        if rank is not None and rank >= STEP2_MIN_RANK_ID:
            survivors.append(row)
    return survivors


def _normalize_stray_quote_wrapped_value(value: str | None) -> str | None:
    """A small number of real source-CSV rows carry a literal, stray extra pair of double-quote
    characters wrapped around every field's value (an apparent upstream data-serialization
    artifact) -- e.g. an `r200_nm` value of the literal text `"Magnoliophyta"` (quote characters
    included) rather than `Magnoliophyta`. FR-QPB-118's Step 3 text states only a literal, exact
    seven-value match, with no normalization rule for this artifact; however, Decision Log D-48's
    own cited Step 3 real-data ground truth (26,763 -> 21,025) is only reproducible if these rows
    are still counted as matching their intended phylum. This strips one such wrapping layer
    before the exact-match comparison; the stripped value is what is carried forward into the
    surviving row."""
    if value and len(value) >= 2 and value[0] == '"' and value[-1] == '"':
        return value[1:-1]
    return value


def _step3_vascular_phylum_allowlist_filter(rows: list[dict]) -> list[dict]:
    """FR-QPB-118 Step 3 (Decision Log D-48): keep only rows whose `r200_nm` is exactly one of
    `VASCULAR_PHYLUM_ALLOWLIST`'s seven phyla. A missing/blank `r200_nm` excludes the row."""
    survivors = []
    for row in rows:
        normalized = _normalize_stray_quote_wrapped_value(row.get("r200_nm"))
        if normalized in VASCULAR_PHYLUM_ALLOWLIST:
            if normalized != row.get("r200_nm"):
                row = {**row, "r200_nm": normalized}
            survivors.append(row)
    return survivors


def _load_nibr_vascular_name_lookup(xlsx_path: str | Path) -> dict[str, set[str]]:
    """FR-QPB-118 Step 4 (Decision Log D-48): loads a normalized-scientific-name -> set-of-KTSN
    lookup from the NIBR xlsx workbook's `관속식물류` sheet (`학명`/`KTSN` columns, data starting
    at row 3), for the duplicate-scientific-name canonical-KTSN cross-reference. Raises
    `ReferenceDataMissingError`/`ReferenceDataInvalidError` (mirroring FR-QPB-105/FR-QPB-113's
    existing fail-early discipline for this same workbook) if the workbook, its `관속식물류` sheet,
    or that sheet's `학명`/`KTSN` header cells are missing/malformed."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.is_file():
        raise ReferenceDataMissingError(
            f"식별 기능에 필요한 국가생물종목록 xlsx 파일을 찾을 수 없습니다: {xlsx_path}. "
            "'storage/reference/tables/2025년 국가생물종목록_v1.0.xlsx' 파일이 올바른 위치에 "
            "있는지 확인한 뒤 다시 시도해 주세요."
        )

    import openpyxl

    try:
        workbook = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- openpyxl raises assorted errors for corrupt files
        raise ReferenceDataInvalidError(
            f"국가생물종목록 xlsx 파일을 읽는 중 오류가 발생했습니다: {xlsx_path} ({exc})"
        ) from exc

    try:
        if NIBR_VASCULAR_SHEET_NAME not in workbook.sheetnames:
            raise ReferenceDataInvalidError(
                f"국가생물종목록 xlsx 파일에 '{NIBR_VASCULAR_SHEET_NAME}' 시트가 없습니다: "
                f"{xlsx_path}"
            )
        worksheet = workbook[NIBR_VASCULAR_SHEET_NAME]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise ReferenceDataInvalidError(
                f"국가생물종목록 xlsx 파일의 '{NIBR_VASCULAR_SHEET_NAME}' 시트에 헤더 행이 "
                f"없습니다: {xlsx_path}"
            )

        name_col_idx: int | None = None
        ktsn_col_idx: int | None = None
        for idx, cell_value in enumerate(header_row):
            text = str(cell_value).strip() if cell_value is not None else ""
            if text == NIBR_VASCULAR_SHEET_NAME_HEADER and name_col_idx is None:
                name_col_idx = idx
            elif text == NIBR_VASCULAR_SHEET_KTSN_HEADER and ktsn_col_idx is None:
                ktsn_col_idx = idx

        if name_col_idx is None or ktsn_col_idx is None:
            missing = [
                label
                for label, idx in (
                    (NIBR_VASCULAR_SHEET_NAME_HEADER, name_col_idx),
                    (NIBR_VASCULAR_SHEET_KTSN_HEADER, ktsn_col_idx),
                )
                if idx is None
            ]
            raise ReferenceDataInvalidError(
                f"국가생물종목록 xlsx 파일의 '{NIBR_VASCULAR_SHEET_NAME}' 시트에 필수 열이 "
                f"없습니다: {missing} (파일: {xlsx_path})"
            )

        lookup: dict[str, set[str]] = {}
        for row in worksheet.iter_rows(
            min_row=NIBR_VASCULAR_SHEET_DATA_START_ROW, values_only=True
        ):
            if row is None:
                continue
            name_value = row[name_col_idx] if len(row) > name_col_idx else None
            ktsn_value = row[ktsn_col_idx] if len(row) > ktsn_col_idx else None
            if name_value is None or ktsn_value is None:
                continue
            normalized_name = _normalize_taxon_full_nm(str(name_value))
            if not normalized_name:
                continue
            lookup.setdefault(normalized_name, set()).add(str(ktsn_value))
        return lookup
    except (ReferenceDataMissingError, ReferenceDataInvalidError):
        raise
    except Exception as exc:  # noqa: BLE001 -- any other parsing failure must fail early, not silently
        raise ReferenceDataInvalidError(
            f"국가생물종목록 xlsx 파일을 처리하는 중 오류가 발생했습니다: {xlsx_path} ({exc})"
        ) from exc
    finally:
        workbook.close()


def _step4_duplicate_name_resolution(rows: list[dict], xlsx_path: str | Path | None) -> list[dict]:
    """FR-QPB-118 Step 4 (Decision Log D-48): groups rows by normalized `taxon_full_nm`. A group of
    size 1 survives unchanged. For a group of size > 1, resolves against the NIBR xlsx's
    `관속식물류` sheet: if exactly one distinct `KTSN` value is found for the normalized name, and
    it equals one of the group's own `ktsn` values, keeps only that one row; otherwise drops the
    entire group. The xlsx is required (and validated) whenever this step runs."""
    if xlsx_path is None:
        raise ReferenceDataMissingError(
            "중복된 학명을 해소하는 데 필요한 국가생물종목록 xlsx 파일 경로가 제공되지 않았습니다 "
            "(FR-QPB-118 Step 4)."
        )
    lookup = _load_nibr_vascular_name_lookup(xlsx_path)

    groups: dict[str, list[dict]] = {}
    for row in rows:
        name = _normalize_taxon_full_nm(row.get("taxon_full_nm"))
        groups.setdefault(name, []).append(row)

    keep_ids: set[int] = set()
    for name, group in groups.items():
        if len(group) == 1:
            keep_ids.add(id(group[0]))
            continue
        candidates = lookup.get(name, set())
        own_ktsns = {row.get("ktsn") for row in group}
        if len(candidates) == 1:
            (only_value,) = candidates
            if only_value in own_ktsns:
                for row in group:
                    if row.get("ktsn") == only_value:
                        keep_ids.add(id(row))
                        break
        # Every other case (name absent from the sheet; more than one distinct KTSN found; or the
        # one value found matches none of the group's own ktsn values) drops the entire group --
        # never a fallback tie-break of any kind.
    return [row for row in rows if id(row) in keep_ids]


def _step5_project_lookup_columns(rows: list[dict]) -> list[dict]:
    """FR-QPB-118 Step 5 (unchanged in substance from the original, single-step form): project each
    row down to exactly `REQUIRED_KTSN_CSV_COLUMNS`, verbatim (no reformatting/renormalization
    beyond the projection itself)."""
    return [{col: row[col] for col in REQUIRED_KTSN_CSV_COLUMNS} for row in rows]


def _run_pipeline_rows(
    csv_path: str | Path, xlsx_path: str | Path | None, through_step: int
) -> list[dict]:
    rows = _load_full_csv_rows(Path(csv_path))
    if through_step >= 1:
        rows = _step1_plantae_subtree_filter(rows)
    if through_step >= 2:
        rows = _step2_rank_filter(rows)
    if through_step >= 3:
        rows = _step3_vascular_phylum_allowlist_filter(rows)
    if through_step >= 4:
        rows = _step4_duplicate_name_resolution(rows, xlsx_path)
    if through_step >= 5:
        rows = _step5_project_lookup_columns(rows)
    return rows


def run_ktsn_reference_pipeline(
    csv_path: str, xlsx_path: str | None = None, through_step: int = 5
) -> dict:
    """HARNESS_CONTRACT.md function 11 -- a pure-Python reference implementation of FR-QPB-118
    (further revised; Decision Log D-48)'s five ordered row-filtering/extraction steps, exposing
    each step's intermediate output (AC-QPB-084/085/086/087). See HARNESS_CONTRACT.md for the
    authoritative parameter/return-shape contract. This is a thin wrapper over the same
    `_run_pipeline_rows`/`_step*` internals `_extract_ktsn_lookup_csv` uses for the real,
    generated-project artifact -- not a second, independent implementation of the five steps."""
    rows = _run_pipeline_rows(csv_path, xlsx_path, through_step)
    return {"rows": rows, "row_count": len(rows)}


def _step5_accepted_only_filter(rows: list[dict]) -> list[dict]:
    """FR-QPB-126 Step 5 (Decision Log D-67, new -- not the same "Step 5" as FR-QPB-118's own
    unrelated column-projection step above): keep only rows whose `taxon_jm_nm` value is exactly
    `ACCEPTED_TAXON_STATUS_VALUE` ("정명") -- excludes every other row, including one with a
    missing/blank `taxon_jm_nm` value (fails the exact-match test), matching this pipeline's
    existing exclude-on-ambiguity posture (Decision Log D-48)."""
    return [row for row in rows if row.get("taxon_jm_nm") == ACCEPTED_TAXON_STATUS_VALUE]


def derive_accepted_name_lookup_table(csv_path: str, xlsx_path: str | None = None) -> dict:
    """HARNESS_CONTRACT.md function 17 -- FR-QPB-126 (Decision Log D-67)'s own row-derivation rule
    for DR-QPB-072's bundled accepted-name lookup table: (1) reuses the identical, shared Steps 1-4
    row-filtering computation `run_ktsn_reference_pipeline(csv_path, xlsx_path, through_step=4)`
    already exposes (Plantae-kingdom-subtree filter; species-level-and-below rank filter;
    vascular-plants-only phylum allowlist; duplicate-scientific-name canonical-KTSN resolution);
    then (2) applies the new Step 5 accepted-only filter above; then (3) projects each surviving
    row down to exactly `ACCEPTED_NAME_LOOKUP_TABLE_COLUMNS`, with `taxon_full_nm` tag-stripped
    (mirrors FR-QPB-106/107's existing `<em>`/`</em>`-removal rule) and `ktsn` preserved as a
    string.

    Returns `{"rows": [{"ktsn": str, "taxon_kor_nm": str, "taxon_full_nm": str}],
    "row_count": int}`.

    This is a test-only-shaped, pure-Python reference implementation of the *rule* -- mirroring
    this module's existing `run_ktsn_reference_pipeline`/`_extract_ktsn_lookup_csv` "rule in
    Python, artifact via build_project" split -- the real, generated-project artifact (a real,
    indexed GeoPackage table) is bundled separately, by :func:`qfield_builder.gpkg.
    add_ktsn_lookup_table`, called from :mod:`qfield_builder.build` with this same function's own
    `"rows"` output.
    """
    steps_1_to_4 = run_ktsn_reference_pipeline(csv_path, xlsx_path, through_step=4)["rows"]
    accepted_only = _step5_accepted_only_filter(steps_1_to_4)
    rows = [
        {
            "ktsn": row["ktsn"],
            "taxon_kor_nm": row.get("taxon_kor_nm") or "",
            "taxon_full_nm": _normalize_taxon_full_nm(row.get("taxon_full_nm")),
        }
        for row in accepted_only
    ]
    return {"rows": rows, "row_count": len(rows)}


def load_filtered_accepted_name_lookup(reference_data_dir: str | Path) -> dict:
    """Load the accepted-only CSV from a validated D-89 reference bundle."""
    _lookup, accepted_path, _nibr, _raster = _validate_filtered_manifest(Path(reference_data_dir))
    rows: list[dict] = []
    try:
        with accepted_path.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            if reader.fieldnames != list(ACCEPTED_NAME_LOOKUP_TABLE_COLUMNS):
                raise ReferenceDataInvalidError(
                    f"필터링된 인정 국명 lookup의 열이 올바르지 않습니다: {accepted_path}"
                )
            for row in reader:
                rows.append(
                    {column: row.get(column, "") for column in ACCEPTED_NAME_LOOKUP_TABLE_COLUMNS}
                )
    except ReferenceDataInvalidError:
        raise
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ReferenceDataInvalidError(
            f"필터링된 인정 국명 lookup을 읽을 수 없습니다: {accepted_path} ({exc})"
        ) from exc
    return {"rows": rows, "row_count": len(rows)}


def _read_national_ktsn_values(xlsx_path: Path) -> list[str]:
    """Read the strict, one-value-per-line NIBR list used by D-40 and D-89."""
    if not xlsx_path.is_file():
        raise ReferenceDataMissingError(f"국가생물종목록 xlsx 파일을 찾을 수 없습니다: {xlsx_path}")
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ReferenceDataInvalidError(
            f"국가생물종목록 xlsx 파일을 읽을 수 없습니다: {xlsx_path} ({exc})"
        ) from exc
    try:
        if NATIONAL_LIST_XLSX_SHEET_NAME not in workbook.sheetnames:
            raise ReferenceDataInvalidError(
                f"국가생물종목록 xlsx에 '{NATIONAL_LIST_XLSX_SHEET_NAME}' 시트가 없습니다"
            )
        values = []
        worksheet = workbook[NATIONAL_LIST_XLSX_SHEET_NAME]
        for row in worksheet.iter_rows(
            min_row=NATIONAL_LIST_XLSX_DATA_START_ROW, max_col=1, values_only=True
        ):
            value = row[0] if row else None
            if value is not None and str(value).strip():
                values.append(str(value).strip())
        if not values:
            raise ReferenceDataInvalidError(
                f"국가생물종목록 xlsx의 KTSN 목록이 비어 있습니다: {xlsx_path}"
            )
        return values
    finally:
        workbook.close()


def validate_canonical_release_source(
    workbook_path: str | Path, manifest_path: str | Path
) -> dict:
    """Validate the explicitly provisioned canonical workbook against release metadata.

    The manifest is tracked with the release source descriptor; the workbook itself may be
    supplied by a protected artifact store or a local provisioning step.  This keeps the source
    reproducible without treating the workbook as generated-project data.
    """
    workbook = Path(workbook_path)
    metadata_path = Path(manifest_path)
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ReferenceDataInvalidError(
            f"canonical release source manifest를 읽을 수 없습니다: {metadata_path} ({exc})"
        ) from exc
    entry = metadata.get("canonical_workbook") if isinstance(metadata, dict) else None
    if not isinstance(entry, dict):
        raise ReferenceDataInvalidError(
            "canonical release source manifest에 canonical_workbook 항목이 없습니다: "
            f"{metadata_path}"
        )
    if workbook.name != entry.get("filename"):
        raise ReferenceDataInvalidError(
            "provisioned canonical workbook 파일명이 release manifest와 다릅니다: "
            f"{workbook.name}"
        )
    if not workbook.is_file():
        raise ReferenceDataMissingError(
            f"provisioned canonical workbook을 찾을 수 없습니다: {workbook}"
        )
    expected_size = entry.get("byte_size")
    if isinstance(expected_size, bool) or not isinstance(expected_size, int) or expected_size < 1:
        raise ReferenceDataInvalidError(
            f"canonical release source manifest의 byte_size가 올바르지 않습니다: {metadata_path}"
        )
    actual_size = workbook.stat().st_size
    if actual_size != expected_size:
        raise ReferenceDataInvalidError(
            "canonical workbook 크기가 release manifest와 다릅니다: "
            f"{actual_size} != {expected_size}"
        )
    expected_sha256 = entry.get("sha256")
    if not isinstance(expected_sha256, str) or len(expected_sha256) != 64:
        raise ReferenceDataInvalidError(
            f"canonical release source manifest의 sha256가 올바르지 않습니다: {metadata_path}"
        )
    actual_sha256 = _sha256_file(workbook)
    if actual_sha256 != expected_sha256:
        raise ReferenceDataInvalidError(
            "canonical workbook 해시가 release manifest와 다릅니다. "
            "승인된 source artifact를 다시 provision해 주세요."
        )
    return {"path": str(workbook), "sha256": actual_sha256, "byte_size": actual_size}


def prepare_filtered_reference_bundle(
    canonical_workbook_or_root: str,
    destination_dir: str,
    *,
    raster_dir: str | Path | None = None,
    expected_sha256: str | None = None,
) -> Path:
    """Create and validate the D-95 release bundle from the canonical workbook atomically.

    ``destination_dir`` is a staging directory consumed by PyInstaller.  It contains only the
    generated lookup artifacts, the manifest, and byte-for-byte raster copies.  The canonical
    workbook is read only as a build input and is never copied into this staging tree.

    ``canonical_workbook_or_root`` may be the workbook itself (the release-script form) or a
    source root containing ``tables/<canonical filename>`` (the source-checkout form).  In both
    forms, ``raster_dir`` must identify the separately provisioned raster directory when the
    workbook itself is passed.
    """
    candidate = Path(canonical_workbook_or_root)
    if candidate.is_dir():
        source = candidate
        workbook_path = source / CANONICAL_WORKBOOK_RELATIVE_SUBPATH
        resolved_raster_dir = source / RASTER_DIR_RELATIVE_SUBPATH
    else:
        workbook_path = candidate
        resolved_raster_dir = Path(raster_dir) if raster_dir is not None else None
        if resolved_raster_dir is None:
            raise ReferenceDataMissingError(
                "canonical release staging에는 별도로 provision된 raster directory가 필요합니다"
            )
        source = workbook_path.parent.parent
    if not workbook_path.is_file():
        raise ReferenceDataMissingError(
            f"release canonical workbook을 찾을 수 없습니다: {workbook_path}"
        )
    if expected_sha256 is not None and _sha256_file(workbook_path) != expected_sha256:
        raise ReferenceDataInvalidError(
            "canonical workbook 해시가 확인된 release source와 다릅니다. "
            "source를 다시 provision해 주세요."
        )
    canonical = canonical_reference.ingest_canonical_workbook(
        str(workbook_path), source_kind="bundled_candidate"
    )
    if not canonical.get("success"):
        raise ReferenceDataInvalidError(
            canonical.get("error_message") or "canonical workbook 검증 실패"
        )
    _validate_raster_dir(resolved_raster_dir)
    canonical_rows = canonical["rows"]
    lookup_rows = [
        {
            "ktsn": row["ktsn"],
            "taxon_full_nm": row["scientific_name"],
            "taxon_kor_nm": row.get("korean_name") or "",
            "taxon_jm_nm": row["taxon_status"],
            "correct_list": (
                "[]"
                if row["taxon_status"] == ACCEPTED_TAXON_STATUS_VALUE
                else json.dumps(
                    [{"KTSN": row["accepted_ktsn"]}],
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
            ),
        }
        for row in canonical_rows
    ]
    accepted_rows = [
        {
            "ktsn": row["ktsn"],
            "taxon_kor_nm": row.get("korean_name") or "",
            "taxon_full_nm": _normalize_taxon_full_nm(row["scientific_name"]),
        }
        for row in canonical_rows
        if row["taxon_status"] == ACCEPTED_TAXON_STATUS_VALUE
    ]
    nibr_values = [
        row["ktsn"]
        for row in canonical_rows
        if row["taxon_status"] == ACCEPTED_TAXON_STATUS_VALUE
    ]

    destination = Path(destination_dir)
    parent = destination.parent
    parent.mkdir(parents=True, exist_ok=True)
    temporary = parent / f".{destination.name}.tmp"
    shutil.rmtree(temporary, ignore_errors=True)
    try:
        filtered_dest = temporary / FILTERED_REFERENCE_RELPATH
        filtered_dest.mkdir(parents=True)
        lookup_dest = filtered_dest / FILTERED_KTSN_LOOKUP_NAME
        with lookup_dest.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(
                fh, fieldnames=REQUIRED_KTSN_CSV_COLUMNS, lineterminator="\n"
            )
            writer.writeheader()
            # The raw source uses an empty cell for most terminal `정명` rows.  The release
            # artifact has a stricter contract: every transition cell is valid JSON, so an
            # empty source value is represented by the equivalent empty JSON array.
            writer.writerows(
                {
                    **row,
                    "correct_list": row.get("correct_list") or "[]",
                }
                for row in lookup_rows
            )
        accepted_dest = filtered_dest / FILTERED_ACCEPTED_LOOKUP_NAME
        with accepted_dest.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(
                fh, fieldnames=ACCEPTED_NAME_LOOKUP_TABLE_COLUMNS, lineterminator="\n"
            )
            writer.writeheader()
            writer.writerows(accepted_rows)
        (filtered_dest / FILTERED_NIBR_LIST_NAME).write_text(
            "\n".join(nibr_values) + "\n", encoding="utf-8"
        )

        raster_dest = temporary / RASTER_DIR_RELATIVE_SUBPATH
        raster_dest.mkdir(parents=True)
        for tif_path in sorted(resolved_raster_dir.glob("*.tif")):
            shutil.copyfile(tif_path, raster_dest / tif_path.name)

        source_provenance = {
            "input_kind": "canonical_workbook",
            "canonical_workbook": {
                # The provenance path is a stable logical source path, never the absolute path
                # supplied by release CI or a developer's local provisioning directory.
                "path": CANONICAL_WORKBOOK_RELATIVE_SUBPATH.as_posix(),
                "sha256": _sha256_file(workbook_path),
            },
        }

        raster_files = {
            path.name: {
                "path": (RASTER_DIR_RELATIVE_SUBPATH / path.name).as_posix(),
                "size": path.stat().st_size,
                "sha256": _sha256_file(path),
            }
            for path in sorted(raster_dest.glob("*.tif"))
        }

        artifact_paths = [
            filtered_dest / name
            for name in (
                FILTERED_KTSN_LOOKUP_NAME,
                FILTERED_ACCEPTED_LOOKUP_NAME,
                FILTERED_NIBR_LIST_NAME,
            )
        ]
        manifest = {
            "pipeline_revision": CANONICAL_FILTERED_PIPELINE_REVISION,
            "source_provenance": source_provenance,
            "raster_files": raster_files,
            "artifacts": {
                path.name: {
                    "path": f"filtered/{path.name}",
                    "row_count": (
                        sum(1 for _ in path.open(encoding="utf-8-sig")) - 1
                        if path.suffix == ".csv"
                        else len(path.read_text(encoding="utf-8").splitlines())
                    ),
                    "sha256": _sha256_file(path),
                }
                for path in artifact_paths
            },
        }
        (temporary / FILTERED_MANIFEST_NAME).write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        _validate_filtered_manifest(temporary, require_release_contract=True)
        shutil.rmtree(destination, ignore_errors=True)
        os.replace(temporary, destination)
        # The manifest is intentionally copied beside `filtered/` so it is not self-hashed.
        shutil.copyfile(
            destination / FILTERED_MANIFEST_NAME, destination.parent / FILTERED_MANIFEST_NAME
        )
        return destination
    except BaseException:
        shutil.rmtree(temporary, ignore_errors=True)
        raise


def _extract_ktsn_lookup_csv(csv_path: Path, dest_path: Path, xlsx_path: Path) -> None:
    """FR-QPB-118 (further revised; Decision Log D-48): runs the complete five-step pipeline
    (Steps 1-4 row filtering, Step 5 column extraction) over the complete raw KTSN CSV at
    `csv_path`, writing the surviving rows' `REQUIRED_KTSN_CSV_COLUMNS` values to a new, small CSV
    at `dest_path`. `ktsn` values are preserved as strings (the `csv` module never renormalizes
    values as numbers) and `correct_list` values are preserved verbatim, exactly as they appear in
    the source, including when empty.

    Writes to a temporary file first and only replaces `dest_path` once the full pipeline has
    succeeded, so a failure partway through (e.g. an I/O error, or a Step 4 NIBR-xlsx failure)
    never leaves a partial/truncated `dest_path` behind -- must be called only after
    :func:`_validate_ktsn_csv` has already succeeded for the same `csv_path`."""
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = dest_path.with_name(dest_path.name + ".tmp")
    try:
        rows = _run_pipeline_rows(csv_path, xlsx_path, through_step=5)
        with tmp_path.open("w", encoding="utf-8", newline="") as dest_fh:
            writer = csv.writer(dest_fh)
            writer.writerow(REQUIRED_KTSN_CSV_COLUMNS)
            for row in rows:
                writer.writerow([row[col] for col in REQUIRED_KTSN_CSV_COLUMNS])
    except (ReferenceDataMissingError, ReferenceDataInvalidError):
        tmp_path.unlink(missing_ok=True)
        raise
    except (OSError, UnicodeDecodeError, csv.Error, KeyError, ValueError) as exc:
        tmp_path.unlink(missing_ok=True)
        raise ReferenceDataInvalidError(
            f"KTSN 참조 CSV 파일에서 5단계 필터링/추출 파이프라인을 실행하는 중 오류가 "
            f"발생했습니다: {csv_path} ({exc})"
        ) from exc
    except BaseException:
        tmp_path.unlink(missing_ok=True)
        raise
    else:
        tmp_path.replace(dest_path)


def bundle_reference_data(
    reference_data_dir: str,
    project_dir: str,
    *,
    include_probability_rasters: bool = True,
) -> None:
    """Legacy compatibility helper for the pre-D-95 CSV/XLSX source contract.

    FR-QPB-112 (revised; Decision Log D-47): bundles (1) the build-time-extracted, 5-column
    KTSN lookup file (FR-QPB-118, further revised; Decision Log D-48's five-step pipeline) at
    `project_dir/reference/ktsn_lookup.csv` -- never the complete, unmodified raw CSV -- and (2)
    the complete, unmodified probability-raster set (unaffected by D-47/D-48, still bundled exactly
    as before), into `project_dir/reference/...`. Must be called only after
    :func:`validate_reference_data` has already succeeded for the same `reference_data_dir`."""
    csv_path, raster_dir = validate_reference_data(reference_data_dir)
    if is_filtered_reference_data(reference_data_dir):
        base = Path(reference_data_dir)
        _lookup, _accepted, nibr, _raster = _validate_filtered_manifest(base)
        destination = Path(project_dir) / PROJECT_REFERENCE_DIR
        destination.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(csv_path, destination / Path(PROJECT_KTSN_LOOKUP_RELPATH).name)
        shutil.copyfile(nibr, destination / Path(PROJECT_NATIONAL_LIST_RELPATH).name)
        if include_probability_rasters:
            dest_raster_dir = Path(project_dir) / PROJECT_RASTER_DIR_RELPATH
            dest_raster_dir.mkdir(parents=True, exist_ok=True)
            for tif_path in sorted(raster_dir.glob("*.tif")):
                shutil.copyfile(tif_path, dest_raster_dir / tif_path.name)
        return
    xlsx_path = Path(reference_data_dir) / NATIONAL_LIST_XLSX_RELATIVE_SUBPATH

    project_root = Path(project_dir)
    dest_lookup_csv = project_root / PROJECT_KTSN_LOOKUP_RELPATH
    _extract_ktsn_lookup_csv(csv_path, dest_lookup_csv, xlsx_path)

    if include_probability_rasters:
        dest_raster_dir = project_root / PROJECT_RASTER_DIR_RELPATH
        dest_raster_dir.mkdir(parents=True, exist_ok=True)
        for tif_path in sorted(raster_dir.glob("*.tif")):
            shutil.copyfile(tif_path, dest_raster_dir / tif_path.name)


def extract_and_bundle_national_ktsn_list(reference_data_dir: str, project_dir: str) -> bool:
    """Legacy compatibility helper; canonical builds materialize this data from the workbook.

    FR-QPB-113 (Decision Log D-41): extracts only the `KTSN` column (column A) from the NIBR
    accepted-taxon xlsx's `62,604종` sheet into a small, single-purpose, one-value-per-line plain
    text lookup list, bundled at `project_dir/reference/nibr_accepted_ktsn_list.txt`.

    This is a best-effort addition, not a hard build-failure condition like FR-QPB-105's CSV
    check: the xlsx source is a large, separate reference asset not covered by the
    `_test_reference_data_dir` test-seam fixtures (HARNESS_CONTRACT.md), so this silently does
    nothing (returns False) when the source xlsx is absent, rather than failing builds that
    intentionally use a minimal reference-data override for fast testing.

    Returns True iff the lookup-list file was actually written.
    """
    if is_filtered_reference_data(reference_data_dir):
        _lookup, _accepted, nibr, _raster = _validate_filtered_manifest(Path(reference_data_dir))
        destination = Path(project_dir) / PROJECT_NATIONAL_LIST_RELPATH
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(nibr, destination)
        return True

    xlsx_path = Path(reference_data_dir) / NATIONAL_LIST_XLSX_RELATIVE_SUBPATH
    if not xlsx_path.is_file():
        return False

    import openpyxl

    workbook = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    try:
        if NATIONAL_LIST_XLSX_SHEET_NAME not in workbook.sheetnames:
            return False
        worksheet = workbook[NATIONAL_LIST_XLSX_SHEET_NAME]
        ktsn_values = []
        for row in worksheet.iter_rows(
            min_row=NATIONAL_LIST_XLSX_DATA_START_ROW, max_col=1, values_only=True
        ):
            value = row[0]
            if value is not None and str(value).strip():
                ktsn_values.append(str(value).strip())
    finally:
        workbook.close()

    if not ktsn_values:
        return False

    dest_path = Path(project_dir) / PROJECT_NATIONAL_LIST_RELPATH
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    dest_path.write_text("\n".join(ktsn_values) + "\n", encoding="utf-8")
    return True
